import pptx

prs = pptx.Presentation('SIH2026_Criminal_Network_Analysis_System.pptx')
print(f"Slide dimensions: {prs.slide_width} x {prs.slide_height}")

with open('inspect_details.txt', 'w', encoding='utf-8') as out:
    out.write(f"Slide dimensions: {prs.slide_width} x {prs.slide_height}\n\n")
    for i, slide in enumerate(prs.slides):
        out.write(f"====================================================\n")
        out.write(f"=== SLIDE {i+1} (shapes: {len(slide.shapes)}) ===\n")
        out.write(f"====================================================\n")
        for j, shape in enumerate(slide.shapes):
            out.write(f"Shape {j}: name={shape.name}, type={shape.shape_type}, left={shape.left}, top={shape.top}, width={shape.width}, height={shape.height}\n")
            if shape.has_text_frame:
                for p_idx, p in enumerate(shape.text_frame.paragraphs):
                    txt = p.text.strip()
                    if txt:
                        runs_str = "; ".join([f"'{r.text}'(font={r.font.name}, sz={r.font.size.pt if r.font.size else 'None'}, b={r.font.bold})" for r in p.runs])
                        out.write(f"  P{p_idx} [level={p.level}]: {txt}\n")
                        out.write(f"     Runs: {runs_str}\n")
            elif shape.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                img = shape.image
                out.write(f"  PICTURE: filename={img.filename}, content_type={img.content_type}, size={len(img.blob)} bytes\n")
print("Done writing inspect_details.txt")
